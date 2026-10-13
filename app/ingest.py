"""Ingest product data from the Syndigo/GDSN spreadsheet exports.

The export is transposed: row 5 (1-indexed) holds attribute names and each
subsequent row is one trade item, with attributes across the columns.

Run directly to (re)load the database and print a gap report:

    python -m app.ingest
"""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl

from . import db
from .gs1 import GS1Error, normalise_gtin

DESKTOP = Path.home() / "OneDrive - JBS USA Food Company" / "Desktop"

DEFAULT_SOURCES = [
    DESKTOP / "JBS-BeefGDSN-JBSDefaultView-03202026.xlsx",
    DESKTOP / "CFS Product Export.xlsx",
]

# Our page model on the left, the GDSN attribute name on the right.
FIELD_MAP = {
    "jbs_sku": "JBS SKU",
    "brand": "MFG Brand Name",
    "trade_item_desc": "tradeItemDescription",
    "short_description": "Short Description",
    "functional_name": "functionalName",
    "net_weight": "netWeight",
    "net_weight_uom": "netWeight UOM",
    "ingredients": "ingredientStatement",
    "storage": "consumerStorageInstructions",
    "prep_serving": "Preparation Serving Information",
    "marketing_message": "tradeItemMarketingMessage",
    "country_of_origin": "Country of Origin",
    "nutrition_raw": "Nutrition Panel",
    "allergens_raw": "Allergen Related Information",
    "image_ref": "Main Product Image",
    "unit_descriptor": "tradeItemUnitDescriptor",
    "sold_to_consumer": "Item is Sold to the Consumer",
}

# Fields a consumer-facing product page cannot really do without.
PAGE_CRITICAL = [
    "brand",
    "trade_item_desc",
    "net_weight",
    "ingredients",
    "nutrition_raw",
    "allergens_raw",
    "prep_serving",
    "image_ref",
]

HEADER_ROW = 4  # 0-indexed row holding attribute names
FIRST_DATA_ROW = 5


def _clean(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_workbook(path: Path) -> list[dict]:
    """Return one dict per trade item, keyed by our page-model field names."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "Attributes" not in wb.sheetnames:
        raise ValueError(f"{path.name}: no 'Attributes' sheet - unexpected export format")

    rows = list(wb["Attributes"].iter_rows(values_only=True))
    header = [_clean(c) for c in rows[HEADER_ROW]]
    index = {name: i for i, name in enumerate(header) if name}

    missing_attrs = [a for a in FIELD_MAP.values() if a not in index]
    if missing_attrs:
        print(f"  ! {path.name}: attributes not present in export: {missing_attrs}")

    gtin_col = index.get("GTIN")
    if gtin_col is None:
        raise ValueError(f"{path.name}: no GTIN column")

    items = []
    for raw in rows[FIRST_DATA_ROW:]:
        gtin_value = _clean(raw[gtin_col]) if gtin_col < len(raw) else None
        if not gtin_value:
            continue
        try:
            gtin14 = normalise_gtin(gtin_value)
        except GS1Error as exc:
            print(f"  ! skipping row: {exc}")
            continue

        item = {"gtin14": gtin14, "source_file": path.name}
        for field, attr in FIELD_MAP.items():
            col = index.get(attr)
            item[field] = _clean(raw[col]) if col is not None and col < len(raw) else None
        items.append(item)
    return items


def gap_report(items: list[dict]) -> dict[str, int]:
    """Count how many items are missing each page-critical field."""
    gaps = {f: 0 for f in PAGE_CRITICAL}
    for item in items:
        for field in PAGE_CRITICAL:
            if not item.get(field):
                gaps[field] += 1
    return gaps


def ingest(paths: list[Path] | None = None) -> list[dict]:
    db.init_db()
    paths = paths or DEFAULT_SOURCES
    all_items: dict[str, dict] = {}

    for path in paths:
        if not path.exists():
            print(f"  ! not found, skipping: {path}")
            continue
        items = read_workbook(path)
        print(f"  + {path.name}: {len(items)} trade items")
        for item in items:
            # Later files win on conflict; they are the more recent export.
            all_items[item["gtin14"]] = item

    with db.cursor() as conn:
        for item in all_items.values():
            db.upsert_product(conn, item)

    return list(all_items.values())


def print_report(items: list[dict]) -> None:
    total = len(items)
    print(f"\n{'=' * 60}\nINGESTED {total} unique trade items\n{'=' * 60}")

    if not total:
        return

    brands: dict[str, int] = {}
    descriptors: dict[str, int] = {}
    for item in items:
        brands[item.get("brand") or "(none)"] = brands.get(item.get("brand") or "(none)", 0) + 1
        d = item.get("unit_descriptor") or "(none)"
        descriptors[d] = descriptors.get(d, 0) + 1

    print("\nBrands:")
    for brand, count in sorted(brands.items(), key=lambda kv: -kv[1]):
        print(f"  {count:5d}  {brand}")

    print("\nPackaging level:")
    for desc, count in sorted(descriptors.items(), key=lambda kv: -kv[1]):
        print(f"  {count:5d}  {desc}")

    print("\nPage-critical data gaps (items missing the field):")
    for field, missing in sorted(gap_report(items).items(), key=lambda kv: -kv[1]):
        pct = 100 * missing / total
        flag = "  <-- blocks the page" if pct == 100 else ""
        print(f"  {missing:5d} / {total}  ({pct:5.1f}%)  {field}{flag}")

    consumer_units = [i for i in items if (i.get("unit_descriptor") or "").upper() != "CASE"]
    print(f"\nNon-CASE (candidate consumer units): {len(consumer_units)} / {total}")
    if not consumer_units:
        print("  Every item in these exports is a CASE-level foodservice item.")
        print("  Retail-ready consumer GTINs will need a different source.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingest GDSN exports into the catalog DB")
    parser.add_argument("files", nargs="*", type=Path, help="xlsx exports (default: Desktop)")
    args = parser.parse_args()
    items = ingest(args.files or None)
    print_report(items)


if __name__ == "__main__":
    main()
